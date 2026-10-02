"""Lecture des segments à traduire et reconstruction du document dans le même format (PDF ou Word)."""
import re
from dataclasses import dataclass, field

BULLET_RX = re.compile(r"^\s*([•●▪◦–—\-*]|\d+(?:\.\d+)*[.)]|\d+(?:\.\d+)+|[a-z][.)])\s+")
MONO_FONTS = ("courier", "mono", "consol", "menlo", "cascadia", "inconsolata", "lucida console", "source code")
HAS_LETTERS = re.compile(r"[A-Za-zÀ-ÿ]{2,}")


@dataclass
class Segment:
    id: int
    text: str                 # texte à traduire (sans puce)
    page: int = 0             # page (PDF) ou 0 (Word)
    prefix: str = ""          # puce ou numéro conservé tel quel
    meta: dict = field(default_factory=dict)


def _translatable(text):
    return bool(HAS_LETTERS.search(text))


# ==========================================================================
# PDF (PyMuPDF)
# ==========================================================================
class PdfDocument:
    kind = "pdf"

    def __init__(self, path):
        import pymupdf
        self.pymupdf = pymupdf
        self.doc = pymupdf.open(path)
        if self.doc.needs_pass:
            raise ValueError("Le PDF est protégé par un mot de passe.")
        self.segments = []
        self._extract()

    @property
    def page_count(self):
        return self.doc.page_count

    def _pieces(self, line):
        """Découpe une ligne en morceaux séparés par de grands espaces (ex. en-tête gauche / droite)."""
        spans = [s for s in line["spans"] if s["text"]]
        pieces, cur = [], []
        for s in spans:
            if cur and s["bbox"][0] - cur[-1]["bbox"][2] > 3 * max(s["size"], 1):
                pieces.append(cur)
                cur = []
            cur.append(s)
        if cur:
            pieces.append(cur)
        return pieces

    def _extract(self):
        fitz = self.pymupdf
        sid = 0
        scanned = 0
        for pno, page in enumerate(self.doc):
            d = page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE | fitz.TEXT_MEDIABOX_CLIP)
            blocks = [b for b in d["blocks"] if b.get("type") == 0]
            if not blocks:
                scanned += 1
            page_groups = []
            for b in blocks:
                groups = []
                for ln in b["lines"]:
                    if abs(ln["dir"][0]) < 0.99:          # texte tourné : ignoré
                        continue
                    for spans in self._pieces(ln):
                        text = "".join(s["text"] for s in spans).strip()
                        if not text:
                            continue
                        main = max(spans, key=lambda s: len(s["text"].strip()))
                        font = main["font"].lower()
                        mono = any(m in font for m in MONO_FONTS) or bool(main["flags"] & 8)
                        bold = bool(main["flags"] & 16) or "bold" in font or "black" in font
                        italic = bool(main["flags"] & 2) or "italic" in font or "oblique" in font
                        x0 = min(s["bbox"][0] for s in spans); y0 = min(s["bbox"][1] for s in spans)
                        x1 = max(s["bbox"][2] for s in spans); y1 = max(s["bbox"][3] for s in spans)
                        rect = fitz.Rect(x0, y0, x1, y1)
                        style = (round(main["size"], 1), bold, mono)
                        new_item = bool(BULLET_RX.match(text))
                        g = groups[-1] if groups else None
                        if (g and not new_item and g["style"] == style
                                and 0 <= rect.y0 - g["rects"][-1].y1 < 0.9 * main["size"]
                                and rect.x0 < g["rects"][-1].x1):
                            g["lines"].append(text)
                            g["rects"].append(rect)
                        else:
                            groups.append(dict(lines=[text], rects=[rect], style=style, size=main["size"],
                                               color=main["color"], bold=bold, italic=italic, mono=mono,
                                               serif=("times" in font or "serif" in font) and "sans" not in font))
                page_groups.extend(groups)
            # Espace disponible : largeur de la colonne de texte et blanc sous chaque bloc
            right = max((r.x1 for g in page_groups if not g["mono"] for r in g["rects"]), default=page.rect.width)
            bottom = page.rect.height - 36
            for g in page_groups:
                r0 = g["rects"][0]
                rect = fitz.Rect(r0)
                for r in g["rects"][1:]:
                    rect |= r
                wide = rect.width > 0.45 * page.rect.width
                if wide:
                    rect.x1 = max(rect.x1, right)
                below = [h["rects"][0].y0 for h in page_groups if h is not g and h["rects"][0].y0 >= rect.y1 - 1
                         and h["rects"][0].x0 < rect.x1 and h["rects"][-1].x1 > rect.x0]
                limit = min(below + [bottom]) - 1.5
                if wide and limit > rect.y1:
                    rect.y1 = min(limit, rect.y1 + 1.6 * g["size"])
                g["box"] = rect
                rs = g["rects"]
                g["leading"] = (max(1.0, min(1.6, (rs[-1].y0 - rs[0].y0) / (len(rs) - 1) / g["size"]))
                                if len(rs) > 1 else 1.15)
            for g in page_groups:
                if g["mono"]:
                    continue                                    # code : laissé tel quel
                text = ""
                for ln in g["lines"]:
                    if text.endswith("-") and len(text) > 1 and text[-2].isalpha():
                        text += ln
                    else:
                        text = (text + " " + ln).strip()
                m = BULLET_RX.match(text)
                prefix = m.group(0) if m else ""
                body = text[len(prefix):] if prefix else text
                if not _translatable(body):
                    continue
                self.segments.append(Segment(sid, body, pno + 1, prefix, meta=g))
                sid += 1
        self.scanned_pages = scanned

    def apply(self, translations, progress=None, cancel=None):
        fitz = self.pymupdf
        by_page = {}
        for seg in self.segments:
            if seg.id in translations:
                by_page.setdefault(seg.page, []).append(seg)
        n = len(by_page)
        for k, (pno, segs) in enumerate(sorted(by_page.items())):
            if cancel is not None and cancel.is_set():
                from .engines import Cancelled
                raise Cancelled()
            page = self.doc[pno - 1]
            for seg in segs:
                for r in seg.meta["rects"]:
                    page.add_redact_annot(fitz.Rect(r.x0 - 0.5, r.y0 + 0.6, r.x1 + 0.5, r.y1 - 0.6), fill=None)
            page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE, graphics=fitz.PDF_REDACT_LINE_ART_NONE)
            # Taille harmonisée : même style => même taille sur la page (au moins 82 % de l'original)
            fits = {}
            for seg in segs:
                g = seg.meta
                seg.meta["fit"] = self._fit(seg.prefix + translations[seg.id], g["box"], g["size"], g["bold"])
                key = (round(g["size"], 1), g["bold"])
                fits[key] = min(fits.get(key, g["size"]), max(seg.meta["fit"], 0.82 * g["size"]))
            for seg in segs:
                g = seg.meta
                size = min(fits[(round(g["size"], 1), g["bold"])], g["fit"])
                rect = fitz.Rect(g["box"])
                rect.x1 += 1.5
                rect.y1 += 0.25 * g["size"]
                name = self._font(page, g)
                c = g["color"] & 0xFFFFFF
                color = ((c >> 16) / 255, ((c >> 8) & 255) / 255, (c & 255) / 255)
                text = seg.prefix + translations[seg.id]
                lh = g["leading"] if len(g["rects"]) > 1 else 1.15
                for _ in range(25):                  # réduit la police jusqu'à ce que le texte tienne
                    if page.insert_textbox(rect, text, fontname=name, fontsize=size, color=color, lineheight=lh) >= 0:
                        break
                    size *= 0.96
            if progress:
                progress((k + 1) / max(1, n))

    _FONTS = {(False, False, False): "helv", (True, False, False): "hebo", (False, True, False): "heit",
              (True, True, False): "hebi", (False, False, True): "tiro", (True, False, True): "tibo",
              (False, True, True): "tiit", (True, True, True): "tibi"}

    def _font(self, page, g):
        """Police Unicode complète (Nimbus, intégrée par MuPDF), insérée une seule fois dans le document."""
        key = (g["bold"], g["italic"], g["serif"])
        name = "Sally" + self._FONTS[key].capitalize()
        if not hasattr(self, "_font_buffers"):
            self._font_buffers = {}
        if key not in self._font_buffers:
            self._font_buffers[key] = self.pymupdf.Font(self._FONTS[key]).buffer
        page.insert_font(fontname=name, fontbuffer=self._font_buffers[key])
        return name

    def _fit(self, text, rect, size, bold):
        """Plus grande taille (<= size) pour laquelle le texte tient dans le rectangle (estimation)."""
        font = "hebo" if bold else "helv"
        words = text.split()
        s = size
        while s > 0.5 * size:
            lh = s * 1.15
            lines, cur = 1, 0.0
            space = self.pymupdf.get_text_length(" ", font, s)
            for w in words:
                wl = self.pymupdf.get_text_length(w, font, s)
                if cur and cur + space + wl > rect.width - 2:
                    lines += 1
                    cur = wl
                else:
                    cur += (space if cur else 0) + wl
            if lines * lh <= rect.height + 0.2 * s:
                return s
            s *= 0.97
        return s

    def save(self, path):
        self.doc.subset_fonts()
        self.doc.save(path, garbage=3, deflate=True)
        self.doc.close()


# ==========================================================================
# Word (python-docx) : corps, tableaux, zones de texte, en-têtes, pieds de page, notes
# ==========================================================================
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{%s}" % W_NS
CODE_STYLES = ("code", "source", "preformatted", "macro", "console")


class DocxDocument:
    kind = "docx"

    def __init__(self, path):
        import docx
        self.docx = docx
        self.doc = docx.Document(path)
        self.segments = []
        self._paras = {}
        self._blob_parts = []
        self._extract()

    @property
    def page_count(self):
        return 0

    def _roots(self):
        from docx.opc.part import XmlPart
        from lxml import etree
        seen = set()
        for part in self.doc.part.package.iter_parts():
            ct = part.content_type
            if not any(k in ct for k in ("document.main", "header", "footer", "footnotes", "endnotes")):
                continue
            if id(part) in seen:
                continue
            seen.add(id(part))
            if isinstance(part, XmlPart):
                yield part, part.element
            else:
                root = etree.fromstring(part.blob)
                self._blob_parts.append((part, root))
                yield part, root

    def _is_code(self, p):
        ppr = p.find(W + "pPr")
        if ppr is not None:
            st = ppr.find(W + "pStyle")
            if st is not None and any(c in (st.get(W + "val") or "").lower() for c in CODE_STYLES):
                return True
        fonts = [f.get(W + "ascii") or "" for f in p.iter(W + "rFonts")]
        return bool(fonts) and all(any(m in f.lower() for m in MONO_FONTS) for f in fonts)

    @staticmethod
    def _texts(p):
        """Éléments w:t du paragraphe, hors zones de texte imbriquées (traitées comme paragraphes à part)."""
        out = []
        for t in p.iter(W + "t"):
            anc = t.getparent()
            nested = False
            while anc is not None and anc is not p:
                if anc.tag == W + "txbxContent":
                    nested = True
                    break
                anc = anc.getparent()
            if not nested:
                out.append(t)
        return out

    @staticmethod
    def _part_label(ct):
        for key, label in (("header", "En-tête"), ("footer", "Pied de page"), ("footnotes", "Notes de bas de page"),
                           ("endnotes", "Notes de fin")):
            if key in ct:
                return label
        return "Corps"

    def _extract(self):
        sid = 0
        counters = {}
        for part, root in self._roots():
            label = self._part_label(part.content_type)
            for p in root.iter(W + "p"):
                ts = self._texts(p)
                text = "".join(t.text or "" for t in ts)
                if text.strip():
                    counters[label] = counters.get(label, 0) + 1     # numéro du paragraphe non vide
                if not text.strip() or not _translatable(text) or self._is_code(p):
                    continue
                m = BULLET_RX.match(text)
                prefix = m.group(0) if m and len(m.group(0)) < len(text) else ""
                self.segments.append(Segment(sid, text[len(prefix):].strip(), 0, prefix,
                                             meta={"part": label, "para": counters[label]}))
                self._paras[sid] = (p, ts)
                sid += 1

    def apply(self, translations, progress=None, cancel=None):
        items = list(self._paras.items())
        for k, (sid, (p, ts)) in enumerate(items):
            if sid not in translations:
                continue
            seg = self.segments[sid]
            ts[0].text = seg.prefix + translations[sid]
            ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            for t in ts[1:]:
                t.getparent().remove(t)
            if progress and k % 200 == 0:
                progress(k / max(1, len(items)))
        from lxml import etree
        for part, root in self._blob_parts:
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        if progress:
            progress(1.0)

    def save(self, path):
        self.doc.save(path)


def open_document(path):
    ext = str(path).lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        return PdfDocument(path)
    if ext == "docx":
        return DocxDocument(path)
    if ext == "doc":
        raise ValueError("Format .doc (Word 97-2003) non pris en charge : enregistrez le fichier en .docx.")
    raise ValueError("Format non pris en charge : %s (PDF ou DOCX attendu)." % ext)
